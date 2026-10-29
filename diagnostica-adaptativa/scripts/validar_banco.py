"""Valida banco/itens.json contra banco/matriz.json e gera banco/Banco_Diagnostica.xlsx para revisão.
Uso: python3 validar_banco.py"""
import json, os, re, collections
BASE = os.path.join(os.path.dirname(__file__), "..")
itens = json.load(open(f"{BASE}/banco/itens.json"))["itens"]
mat = json.load(open(f"{BASE}/banco/matriz.json"))
H = mat["habilidades"]
probs = []

def dollars_ok(s):
    return len(re.findall(r"(?<!\\)\$", s)) % 2 == 0

ids = collections.Counter(i["id"] for i in itens)
probs += [f"id duplicado: {k}" for k, v in ids.items() if v > 1]
for i in itens:
    p = i["id"]
    if i["habilidade"] not in H: probs.append(f"{p}: habilidade inexistente")
    elif H[i["habilidade"]]["eixo"] != i["eixo"]: probs.append(f"{p}: eixo não bate com a habilidade")
    if i["nivel"] not in (1, 2, 3, 4): probs.append(f"{p}: nível inválido")
    letras = [a["letra"] for a in i["alternativas"]]
    if letras != list("ABCDE"): probs.append(f"{p}: alternativas fora do padrão A-E")
    if i["gabarito"] not in letras: probs.append(f"{p}: gabarito inválido")
    textos = [a["texto"].strip() for a in i["alternativas"]]
    if len(set(textos)) < 5: probs.append(f"{p}: alternativas repetidas")
    for a in i["alternativas"]:
        if a["letra"] == i["gabarito"] and a["erro"]: probs.append(f"{p}: alternativa correta com erro mapeado")
        if a["letra"] != i["gabarito"] and not a["erro"]: probs.append(f"{p}: distrator {a['letra']} sem erro mapeado")
        if not dollars_ok(a["texto"]): probs.append(f"{p}: $ desbalanceado na alternativa {a['letra']}")
    if not dollars_ok(i["enunciado"]): probs.append(f"{p}: $ desbalanceado no enunciado")
    for f in i["figuras"]:
        if not os.path.exists(f"{BASE}/{f}"): probs.append(f"{p}: figura ausente {f}")
    if not i["resolucao"]: probs.append(f"{p}: sem resolução")

print(f"{len(itens)} itens | problemas: {len(probs)}")
for x in probs: print("  -", x)
print("Origem:", dict(collections.Counter(i["origem"]["tipo"] for i in itens)))
print("Gabaritos:", dict(sorted(collections.Counter(i["gabarito"] for i in itens).items())))
cov = collections.defaultdict(lambda: [0, 0, 0, 0])
for i in itens: cov[i["habilidade"]][i["nivel"] - 1] += 1
print("Cobertura (N1 N2 N3 N4):")
for h in H: print(f"  {h} {cov[h]}  {H[h]['nome'][:60]}")
sem = [h for h in H if sum(cov[h]) == 0]
print("Habilidades sem item:", sem or "nenhuma")

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    wb = Workbook(); ws = wb.active; ws.title = "Itens"
    cab = ["ID", "Origem", "Eixo", "Habilidade", "Nome da habilidade", "Nível", "Gabarito", "Enunciado",
           "A", "B", "C", "D", "E", "Erros mapeados", "Resolução", "Figuras", "Observação"]
    ws.append(cab)
    for c in ws[1]: c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79")
    for i in itens:
        alts = {a["letra"]: a["texto"] for a in i["alternativas"]}
        erros = " | ".join(f"{a['letra']}: {a['erro']}" for a in i["alternativas"] if a["erro"])
        ws.append([i["id"], i["origem"]["prova"], i["eixo"], i["habilidade"], H[i["habilidade"]]["nome"], i["nivel"],
                   i["gabarito"], i["enunciado"], *[alts[l] for l in "ABCDE"], erros, i["resolucao"],
                   ", ".join(i["figuras"]), i["origem"]["nota"]])
    larg = [16, 18, 6, 10, 34, 6, 9, 70, 22, 22, 22, 22, 22, 60, 50, 28, 30]
    for k, w in enumerate(larg): ws.column_dimensions[chr(65 + k)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    m = wb.create_sheet("Matriz")
    m.append(["Código", "Eixo", "Habilidade", "Pré-requisitos", "Peso ENEM", "Peso vestibular", "N1", "N2", "N3", "N4"])
    for c in m[1]: c.font = Font(bold=True)
    for h, v in H.items():
        m.append([h, mat["eixos"][v["eixo"]], v["nome"], ", ".join(v["pre"]), v["peso_enem"], v["peso_vest"], *cov[h]])
    for k, w in enumerate([8, 26, 60, 14, 10, 14, 5, 5, 5, 5]): m.column_dimensions[chr(65 + k)].width = w
    wb.save(f"{BASE}/banco/Banco_Diagnostica.xlsx")
    print("Planilha de revisão gerada: banco/Banco_Diagnostica.xlsx")
except ImportError:
    print("openpyxl ausente: planilha não gerada")
